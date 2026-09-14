"""Acceptance cases use generated, fictional statements; no client data."""
from decimal import Decimal

import pytest
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen.canvas import Canvas

from app.extractors.text_align import extract_text_align
from app.validators import validate_statement

pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))


def statement_pdf(path, pages=None, *, separate=False, date_header="日期", amount_header="借方/贷方金额",
                  controls=True, declarations="", stagger=False, serial=False, missing_header=None,
                  centered_summary=False):
    if pages is None:
        pages = [[("20260101", "收入", "100.00", "100.00", "甲公司"),
                  ("20260102", "支出", "-25.00", "75.00", "乙公司")]]
    canvas = Canvas(str(path), pagesize=(700, 550))
    previous = "0.00"
    for page_no, rows in enumerate(pages, 1):
        canvas.setFont("STSong-Light", 9)
        canvas.drawString(30, 515, "测试银行股份有限公司")
        canvas.drawString(30, 497, "账号: 0001234 货币: 人民币")
        canvas.drawString(30, 479, "账单所属期间: 20260101 20260131")
        if controls:
            canvas.drawString(30, 461, "上页余额: " + previous)
        if page_no != missing_header:
            canvas.drawString(30, 430, date_header)
            canvas.drawString(120, 430, "交易摘要")
            if separate:
                canvas.drawRightString(300, 430, "支出金额")
                canvas.drawRightString(390, 430, "收入金额")
            else:
                canvas.drawRightString(390, 433.8 if stagger else 430, amount_header)
            canvas.drawRightString(480, 430, "账户余额")
            canvas.drawString(481.5, 430, "对方户名")
            if serial:
                canvas.drawString(610, 430, "流水号")
        y = 403
        for date, summary, amount, balance, party in rows:
            canvas.drawString(30, y, date)
            summary_lines = summary.split("\n")
            for offset, part in enumerate(summary_lines):
                canvas.drawString(120, y - offset * 10 + (5 if centered_summary and len(summary_lines) > 1 else 0), part)
            if separate:
                if amount.startswith("-"):
                    canvas.drawRightString(300, y, amount[1:])
                else:
                    canvas.drawRightString(390, y, amount)
            else:
                canvas.drawRightString(390, y, amount)
            canvas.drawRightString(480, y, balance)
            canvas.drawString(481.5, y, party)
            y -= 30
            previous = balance
        if controls:
            canvas.drawString(30, 70, "期末余额: 75.00")
        if declarations:
            canvas.drawString(30, 50, declarations)
        canvas.showPage()
    canvas.save()
    return path


@pytest.mark.parametrize("separate,stagger,date_header", [(False, False, "日期"), (False, True, "交易日期"), (True, False, "交易时间")])
def test_generic_layout_and_evidence(tmp_path, separate, stagger, date_header):
    e = extract_text_align(statement_pdf(tmp_path / "table.pdf", separate=separate, stagger=stagger, date_header=date_header))
    assert e.text_rows_complete, e.errors
    assert e.account_number == "0001234"
    assert e.currency == "CNY"
    assert e.statement_start.isoformat() == "2026-01-01"
    assert e.statement_end.isoformat() == "2026-01-31"
    assert [(t.debit, t.credit, t.balance, t.counterparty) for t in e.transactions] == [
        (Decimal(0), Decimal(100), Decimal(100), "甲公司"),
        (Decimal(25), Decimal(0), Decimal(75), "乙公司"),
    ]
    assert all(t.serial_number == "" for t in e.transactions)
    assert [t.source_row for t in e.transactions] == [1, 2]
    report = validate_statement(e)
    assert report.status == "通过", report.errors
    assert "全页" in report.validation_basis
    assert e.declared_count is e.declared_debit is e.declared_credit is None
    assert len([w for w in report.warnings if "该项目未检查" in w]) == 3


def test_repeated_headers_page_carry_and_wrapped_summary(tmp_path):
    pages = [[("20260101", "收入\n附言", "100.00", "100.00", "甲公司")],
             [("20260102", "支出", "-25.00", "75.00", "乙公司")]]
    e = extract_text_align(statement_pdf(tmp_path / "multi.pdf", pages=pages))
    assert validate_statement(e).status == "通过", e.errors
    assert [t.source_page for t in e.transactions] == [1, 2]
    assert e.transactions[0].summary == "收入\n附言"
    assert e.transactions[1].summary == "支出"
    assert e.opening_balance == 0 and e.closing_balance == 75


def test_wrapped_summary_above_date_belongs_to_current_row(tmp_path):
    rows = [("20260101", "收入", "100.00", "100.00", "甲公司"),
            ("20260102", "本笔第一行\n本笔第二行", "-25.00", "75.00", "乙公司")]
    e = extract_text_align(statement_pdf(tmp_path / "centered.pdf", pages=[rows], centered_summary=True))
    assert validate_statement(e).status == "通过", e.errors
    assert [t.summary for t in e.transactions] == ["收入", "本笔第一行\n本笔第二行"]


@pytest.mark.parametrize("field,bad_value", [(0, "20261301"), (2, "-XX.00"), (2, ""), (3, "7"), (3, "NaN"), (3, "76.00")])
def test_bad_transaction_cannot_pass(tmp_path, field, bad_value):
    row = ["20260102", "支出", "-25.00", "75.00", "乙公司"]
    row[field] = bad_value
    e = extract_text_align(statement_pdf(tmp_path / "bad.pdf", pages=[[
        ("20260101", "收入", "100.00", "100.00", "甲公司"), tuple(row)]]))
    assert validate_statement(e).status == "核对失败"


def test_missing_header_on_any_page_is_not_ignored(tmp_path):
    pages = [[("20260101", "收入", "100.00", "100.00", "甲公司")],
             [("20260102", "支出", "-25.00", "75.00", "乙公司")]]
    e = extract_text_align(statement_pdf(tmp_path / "missing.pdf", pages=pages, missing_header=2))
    assert not e.text_rows_complete
    assert any("第2页" in error for error in e.errors)
    assert validate_statement(e).status == "核对失败"


def test_missing_controls_cannot_use_computed_totals(tmp_path):
    e = extract_text_align(statement_pdf(tmp_path / "no-controls.pdf", controls=False))
    assert e.text_rows_complete
    assert validate_statement(e).status == "核对失败"


def test_existing_declaration_mismatch_is_never_waived(tmp_path):
    e = extract_text_align(statement_pdf(tmp_path / "wrong-total.pdf", declarations="交易笔数: 3 借方合计: 25.00 贷方合计: 100.00"))
    report = validate_statement(e)
    assert e.declared_count == 3
    assert report.status == "核对失败"
    assert any("声明笔数=3" in error for error in report.errors)


@pytest.mark.parametrize("declaration", ["借方合计: 26.00", "贷方合计: 101.00"])
def test_partial_bad_declaration_cannot_use_balance_alternative(tmp_path, declaration):
    e = extract_text_align(statement_pdf(tmp_path / "partial-bad.pdf", declarations=declaration))
    assert validate_statement(e).status == "核对失败"


def test_serial_column_present_but_empty_is_still_error(tmp_path):
    e = extract_text_align(statement_pdf(tmp_path / "serial.pdf", serial=True))
    assert e.serial_number_required
    assert any("流水号为空" in error for error in validate_statement(e).errors)


def test_unsigned_ambiguous_amount_header_is_not_assumed(tmp_path):
    assert extract_text_align(statement_pdf(tmp_path / "ambiguous.pdf", amount_header="金额")) is None


def test_empty_page_is_not_auto_approved(tmp_path):
    e = extract_text_align(statement_pdf(tmp_path / "empty.pdf", pages=[[]]))
    assert validate_statement(e).status == "核对失败"


def test_same_date_same_amount_rows_are_not_deduplicated(tmp_path):
    rows = [("20260101", "收入", "100.00", "100.00", "甲公司"),
            ("20260101", "付款", "-12.50", "87.50", "乙公司"),
            ("20260101", "付款", "-12.50", "75.00", "乙公司")]
    e = extract_text_align(statement_pdf(tmp_path / "duplicates.pdf", pages=[rows]))
    assert len(e.transactions) == 3
    assert validate_statement(e).status == "通过"


def test_summary_containing_footer_label_does_not_truncate_rows(tmp_path):
    rows = [("20260101", "期末余额调整", "100.00", "100.00", "甲公司"),
            ("20260102", "支出", "-25.00", "75.00", "乙公司")]
    e = extract_text_align(statement_pdf(tmp_path / "footer-word.pdf", pages=[rows]))
    assert len(e.transactions) == 2
    assert validate_statement(e).status == "通过"


def test_summary_bank_name_is_not_the_issuing_bank(tmp_path):
    rows = [("20260101", "网上企业银行手续费", "100.00", "100.00", "其他银行"),
            ("20260102", "支出", "-25.00", "75.00", "乙公司")]
    e = extract_text_align(statement_pdf(tmp_path / "issuer.pdf", pages=[rows]))
    assert e.bank_name == "测试银行"


def test_transaction_text_cannot_supply_independent_controls(tmp_path):
    from app.routing import process_pdf
    rows = [("20260101", "期初余额: 0.00", "100.00", "100.00", "甲公司"),
            ("20260102", "期末余额: 75.00", "-25.00", "75.00", "乙公司")]
    result = process_pdf(statement_pdf(tmp_path / "false-controls.pdf", pages=[rows], controls=False), None)
    assert result.extraction.opening_balance is None
    assert result.extraction.closing_balance is None
    assert result.validation.status == "核对失败"


def test_local_success_avoids_model_and_export_shows_basis(tmp_path, monkeypatch):
    from app import routing
    from app.config import ModelSettings
    from app.services import process_batch
    from app.exporters import write_batch_workbook
    from openpyxl import load_workbook
    def unexpected_model(*args, **kwargs):
        pytest.fail("Valid local statements must not call the model")
    monkeypatch.setattr(routing, "extract_ai", unexpected_model)
    path = statement_pdf(tmp_path / "local.pdf")
    events = []
    batch = process_batch([path], ModelSettings("https://example.invalid/v1", "test"), False,
                          progress=lambda *event: events.append(event))
    assert batch.statements[0].validation.status == "通过"
    assert any("第 1/1 页" in message for _, _, message in events)
    book = load_workbook(write_batch_workbook(batch, tmp_path))
    values = list(book["提取汇总"].values)
    row = dict(zip(*values))
    assert row["声明笔数"] is None
    assert "全页" in row["核对依据"]
    assert "未检查" in row["核对提示"]
    assert book["流水明细"].max_row == 3
    book.close()


def test_failed_rows_are_exported_for_review(tmp_path):
    from app.models import BatchResult, ProcessedStatement
    from app.exporters import write_batch_workbook
    from openpyxl import load_workbook
    e = extract_text_align(statement_pdf(tmp_path / "no-evidence.pdf", controls=False))
    report = validate_statement(e)
    book = load_workbook(write_batch_workbook(BatchResult([ProcessedStatement(e.source_path, e, report)]), tmp_path))
    assert book["流水明细"].max_row == 1
    assert book["待核查流水"].max_row == 3
    book.close()
