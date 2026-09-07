from copy import deepcopy
from datetime import date
from decimal import Decimal
from pathlib import Path

from app.models import BatchResult, ProcessedStatement
from app.validators import validate_statement


def control(extraction):
    extraction.opening_balance = Decimal("0")
    extraction.closing_balance = Decimal("75")
    extraction.opening_page = 1
    extraction.closing_page = 1
    extraction.currency = "CNY"
    extraction.statement_start = date(2026, 1, 1)
    extraction.statement_end = date(2026, 1, 31)
    return extraction


def test_整体错移不能只靠连续余额通过(sample_extraction):
    e = control(sample_extraction)
    for row in e.transactions:
        row.balance += Decimal("900")
    report = validate_statement(e)
    assert report.status == "核对失败"
    assert not report.balance_break_rows
    assert report.control_checks["首笔余额"]["difference"] == Decimal("900")
    assert report.control_checks["末笔余额"]["difference"] == Decimal("900")


def test_缺控制余额不冒充通过(sample_extraction):
    report = validate_statement(sample_extraction)
    assert all(c["status"] == "未检查" for c in report.control_checks.values())


def test_单笔及零交易核对(sample_extraction, zero_extraction):
    e = control(sample_extraction)
    e.transactions = e.transactions[:1]
    e.declared_count = 1
    e.declared_debit = Decimal(0)
    e.closing_balance = Decimal(100)
    assert validate_statement(e).status == "通过"
    z = control(zero_extraction)
    assert validate_statement(z).status == "核对失败"
    z.closing_balance = Decimal(0)
    assert validate_statement(z).status == "零交易通过"


def test_文本控制信息和账期不能靠交易推算(sample_extraction):
    from app.对账单控制信息 import enrich_from_pages
    e = enrich_from_pages(sample_extraction, [
        "币种 人民币(CNY) 期初余额 0.00 起始日期20260101 截止日期20260131",
        "期末余额 75.00",
    ])
    assert e.opening_balance == 0 and e.closing_balance == 75
    assert e.opening_page == 1 and e.closing_page == 2
    assert e.statement_start == date(2026, 1, 1)
    assert e.statement_end == date(2026, 1, 31)
    assert e.currency == "CNY"
    other = deepcopy(e)
    other.opening_balance = None
    other.closing_balance = None
    other.statement_start = other.statement_end = None
    enrich_from_pages(other, ["交易日期 20260102 余额 75.00"])
    assert other.statement_start is None and other.opening_balance is None


def test_期间衔接_缺期重叠跨币种(sample_extraction):
    from app.对账单控制信息 import statement_links
    a = control(sample_extraction)
    b = deepcopy(a)
    b.source_path = Path("二月.pdf")
    b.statement_start, b.statement_end = date(2026, 2, 1), date(2026, 2, 28)
    b.opening_balance = Decimal(75)
    def rows():
        return statement_links(BatchResult([ProcessedStatement(e.source_path, e, validate_statement(e)) for e in [a,b]]))
    assert rows()[0]["检查结果"] == "通过"
    b.opening_balance = Decimal(76)
    assert rows()[0]["余额差额"] == 1
    b.statement_start = date(2026, 3, 1)
    b.statement_end = date(2026, 3, 31)
    assert "缺期" in rows()[0]["检查结果"]
    assert rows()[0]["余额差额"] is None
    b.statement_start = date(2026, 1, 15)
    assert "重叠" in rows()[0]["检查结果"]
    b.currency = "USD"
    assert all(r["后份文件"] == "" for r in rows())


def test_回退不能丢失控制余额(sample_extraction, monkeypatch):
    from app import routing
    e = control(sample_extraction)
    for row in e.transactions:
        row.balance += Decimal(900)
    fallback = deepcopy(e)
    fallback.opening_balance = fallback.closing_balance = None
    monkeypatch.setattr(routing, "iter_local_extractors", lambda: [("一", lambda _: e), ("二", lambda _: fallback)])
    result = routing.process_pdf(Path("不存在的合成.pdf"), None)
    assert result.validation.status == "核对失败"
    assert result.extraction.opening_balance == 0


def test_输出保留控制差额和待核查流水(sample_extraction, tmp_path):
    from app.exporters import write_batch_workbook
    from openpyxl import load_workbook
    e = control(sample_extraction)
    e.transactions[0].balance += 900
    e.transactions[1].balance += 900
    batch = BatchResult([ProcessedStatement(e.source_path, e, validate_statement(e))])
    path = write_batch_workbook(batch, tmp_path)
    book = load_workbook(path)
    assert "期间衔接检查" in book.sheetnames
    assert "待核查流水" in book.sheetnames
    assert book["待核查流水"].max_row == 3
    values = list(book["提取汇总"].values)
    assert values[1][values[0].index("首笔余额差额")] == 900
    book.close()


def test_农商行月份不能直接冒充账期但应保留币种():
    from test_chengdu_rural import _parse
    e=_parse()
    assert e.currency=="CNY"
    assert e.statement_start is None
    assert e.statement_end is None


def test_衔接嵌套期间与遮盖账号(sample_extraction):
    from app.对账单控制信息 import statement_links
    a=control(sample_extraction)
    b=deepcopy(a);b.source_path=Path("短账期.pdf")
    b.statement_start=date(2026,1,5);b.statement_end=date(2026,1,10)
    c=deepcopy(a);c.source_path=Path("二月.pdf")
    c.statement_start=date(2026,2,1);c.statement_end=date(2026,2,28)
    c.opening_balance=Decimal(75)
    batch=BatchResult([ProcessedStatement(e.source_path,e,validate_statement(e)) for e in [a,b,c]])
    rows=statement_links(batch)
    assert "重叠" in rows[0]["检查结果"]
    assert rows[1]["检查结果"]=="通过"
    c.account_number="123***456"
    assert any("完整账号" in row["检查结果"] for row in statement_links(batch))



def test_桌面入口导出新的检查表(tmp_path, monkeypatch):
    import time
    from app import ui
    from app.config import AppConfig, ModelSettings
    from openpyxl import load_workbook
    test_真实PDF通过批量入口读控制信息(tmp_path)
    monkeypatch.setattr(ui,"load_app_config",lambda:AppConfig())
    monkeypatch.setattr(ui.PDFExtractorApp,"_load_saved_key",lambda self:"")
    monkeypatch.setattr(ui.PDFExtractorApp,"_save_settings",lambda self:ModelSettings("",""))
    monkeypatch.setattr(ui.filedialog,"askopenfilenames",lambda **_: [str(tmp_path/"中行零交易.pdf")])
    messages=[]
    monkeypatch.setattr(ui.messagebox,"showinfo",lambda *args:messages.append(args))
    monkeypatch.setattr(ui.messagebox,"showerror",lambda *args:messages.append(args))
    app=ui.PDFExtractorApp()
    deadline=time.monotonic()+20
    def poll():
        if messages or time.monotonic()>deadline:
            app.root.quit()
        else:
            app.root.after(50,poll)
    def start():
        app._choose_files()
        app.start_button.invoke()
        app.root.after(50,poll)
    app.root.after(50,start)
    try:
        app.run()
        assert messages and messages[0][0]=="处理完成",messages
        paths=list((tmp_path/"提取结果").glob("*.xlsx"))
        assert len(paths)==1
        book=load_workbook(paths[0])
        assert "期间衔接检查" in book.sheetnames
        rows=list(book["提取汇总"].values)
        assert rows[1][rows[0].index("原件期初余额")]==100
        book.close()
    finally:
        app.root.destroy()


def test_起止日期范围可以用于衔接(sample_extraction):
    from app.对账单控制信息 import enrich_from_pages
    e=enrich_from_pages(sample_extraction,["起止日期：2026/01/01 - 2026/01/31 币种：人民币"])
    assert e.statement_start==date(2026,1,1)
    assert e.statement_end==date(2026,1,31)


def test_真实PDF通过批量入口读控制信息(tmp_path):
    from reportlab.pdfgen import canvas
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from app.services import process_batch
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    path=tmp_path/"中行零交易.pdf"
    c=canvas.Canvas(str(path),pagesize=(1200,800))
    c.setFont("STSong-Light",12)
    for y,text in enumerate([
        "中国银行 Account No. 账号 123456789 起始日期20260101 截止日期20260131",
        "币种 人民币(CNY) 承前页余额 100.00",
        "借方合计 0.00 贷方合计 0.00 本页余额 100.00 本对账期末余额 100.00",
    ]):
        c.drawString(30,760-y*30,text)
    c.save()
    batch=process_batch([path],None,False)
    e=batch.statements[0].extraction
    assert e is not None
    assert e.opening_balance==Decimal("100")
    assert e.statement_start==date(2026,1,1)
    assert batch.statements[0].validation.status=="零交易通过"
