from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from app.exporters import write_batch_workbook
from app.models import BatchResult, ProcessedStatement, StatementExtraction, Transaction, ValidationReport


FIXED_TIME = datetime(2026, 8, 25, 19, 30, 45)


def make_transaction(source_file: str) -> Transaction:
    return Transaction(
        occurred_at=datetime(2026, 1, 2, 9, 30),
        summary="测试收款",
        voucher_type="网银",
        voucher_number="00000001",
        debit=Decimal("0.00"),
        credit=Decimal("100.00"),
        balance=Decimal("100.00"),
        serial_number="12345678901234567890",
        counterparty="测试公司/0000123",
        counterparty_bank="测试银行",
        bank_name="中国民生银行",
        account_number="00001234",
        period_start=date(2026, 1, 1),
        period_end=date(2026, 1, 31),
        source_file=source_file,
        source_page=1,
        extraction_method="民生银行固定栏目",
    )


def make_statement(path: Path, status: str) -> ProcessedStatement:
    transaction = make_transaction(path.name)
    extraction = StatementExtraction(
        source_path=path,
        bank_name="中国民生银行",
        account_number="00001234",
        period_start=date(2026, 1, 1),
        period_end=date(2026, 1, 31),
        page_count=1,
        extraction_method="民生银行固定栏目",
        transactions=[transaction],
        declared_count=1,
        declared_debit=Decimal("0.00"),
        declared_credit=Decimal("100.00"),
    )
    return ProcessedStatement(
        source_path=path,
        extraction=extraction,
        validation=ValidationReport(
            status=status,
            actual_count=1,
            actual_debit=Decimal("0.00"),
            actual_credit=Decimal("100.00"),
            balance_break_rows=[] if status == "通过" else [1],
            errors=[] if status == "通过" else ["测试核对失败"],
        ),
    )


def sample_batch() -> BatchResult:
    return BatchResult(
        statements=[
            make_statement(Path("通过.pdf"), "通过"),
            make_statement(Path("失败.pdf"), "核对失败"),
        ]
    )


def test_workbook_contains_detail_and_summary(tmp_path):
    path = write_batch_workbook(sample_batch(), tmp_path, now=FIXED_TIME)
    book = load_workbook(path, data_only=False)

    assert book.sheetnames == ["流水明细", "提取汇总"]
    assert book["流水明细"]["B2"].value == "00001234"
    assert book["流水明细"]["D2"].value == "00000001"
    assert book["流水明细"]["J2"].value == "12345678901234567890"
    assert book["流水明细"]["B2"].number_format == "@"
    assert book["流水明细"]["D2"].number_format == "@"
    assert book["流水明细"]["J2"].number_format == "@"
    assert book["流水明细"].freeze_panes == "A2"
    assert book["流水明细"].max_row == 2


def test_failed_summary_row_is_red_and_excluded_from_detail(tmp_path):
    path = write_batch_workbook(sample_batch(), tmp_path, now=FIXED_TIME)
    book = load_workbook(path)
    summary = book["提取汇总"]
    failed_row = next(row for row in range(2, summary.max_row + 1) if summary.cell(row, 2).value == "核对失败")

    assert summary.cell(failed_row, 1).fill.fgColor.rgb == "00FFC7CE"
    assert summary.cell(failed_row, 1).font.color.rgb == "009C0006"
    assert book["流水明细"].max_row == 2


def test_existing_result_is_not_overwritten(tmp_path):
    first = write_batch_workbook(sample_batch(), tmp_path, now=FIXED_TIME)
    second = write_batch_workbook(sample_batch(), tmp_path, now=FIXED_TIME)

    assert first != second
    assert first.exists()
    assert second.exists()
    assert second.stem.endswith("_v1")


def test_chinese_headers_and_values_have_readable_column_widths(tmp_path):
    path = write_batch_workbook(sample_batch(), tmp_path, now=FIXED_TIME)
    book = load_workbook(path)

    assert book["提取汇总"].column_dimensions["C"].width >= 14
    assert book["提取汇总"].column_dimensions["J"].width >= 14
    assert book["提取汇总"].column_dimensions["O"].width >= 18
