from datetime import datetime
from pathlib import Path
from unicodedata import east_asian_width

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models import BatchResult
from app.对账单控制信息 import LINK_HEADERS, statement_links


DETAIL_HEADERS = [
    "银行名称",
    "账号",
    "交易时间",
    "凭证号码",
    "摘要",
    "凭证类型",
    "借方发生额",
    "贷方发生额",
    "账户余额",
    "流水号",
    "对方户名/账号",
    "对方行名",
    "操作员",
    "交易机构码",
    "对账开始日",
    "对账结束日",
    "来源文件",
    "来源页码",
    "提取方式",
    "来源页内行序",
]
SUMMARY_HEADERS = [
    "来源文件",
    "状态",
    "银行名称",
    "账号",
    "对账开始日",
    "对账结束日",
    "页数",
    "声明笔数",
    "实际笔数",
    "声明借方合计",
    "实际借方合计",
    "声明贷方合计",
    "实际贷方合计",
    "连续余额",
    "提取方式",
    "失败原因",
    "币种", "原件期初余额", "期初来源页", "原件期末余额", "期末来源页",
    "原件账期开始", "原件账期结束", "期初期末核对", "期初期末差额",
    "首笔余额核对", "首笔余额差额", "末笔余额核对", "末笔余额差额", "来源完整路径",
    "核对依据", "核对提示",
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
FAILED_FILL = PatternFill("solid", fgColor="FFC7CE")
FAILED_FONT = Font(color="9C0006")
ZERO_FILL = PatternFill("solid", fgColor="FFF2CC")
ZERO_FONT = Font(color="7F6000")


def _unique_output_path(source_root: Path, now: datetime) -> Path:
    output_dir = Path(source_root) / "提取结果"
    output_dir.mkdir(parents=True, exist_ok=True)
    stem = f"银行流水汇总_{now:%Y%m%d_%H%M%S}"
    candidate = output_dir / f"{stem}.xlsx"
    version = 1
    while candidate.exists():
        candidate = output_dir / f"{stem}_v{version}.xlsx"
        version += 1
    return candidate


def _apply_header(sheet, column_count: int) -> None:
    for cell in sheet[1][:column_count]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(column_count)}{max(sheet.max_row, 1)}"
    sheet.sheet_view.showGridLines = False


def _display_width(value: object) -> int:
    return sum(2 if east_asian_width(character) in {"W", "F", "A"} else 1 for character in str(value or ""))


def _fit_columns(sheet, minimum: int = 10, maximum: int = 40) -> None:
    for column in range(1, sheet.max_column + 1):
        width = max(
            (_display_width(sheet.cell(row, column).value) for row in range(1, sheet.max_row + 1)),
            default=minimum,
        )
        sheet.column_dimensions[get_column_letter(column)].width = min(max(width + 2, minimum), maximum)


def _write_details(sheet, batch: BatchResult, transactions=None) -> None:
    sheet.append(DETAIL_HEADERS)
    for transaction in (batch.approved_transactions if transactions is None else transactions):
        sheet.append(
            [
                transaction.bank_name,
                transaction.account_number,
                transaction.occurred_at,
                transaction.voucher_number,
                transaction.summary,
                transaction.voucher_type,
                transaction.debit,
                transaction.credit,
                transaction.balance,
                transaction.serial_number,
                transaction.counterparty,
                transaction.counterparty_bank,
                transaction.operator,
                transaction.transaction_org_code,
                transaction.period_start,
                transaction.period_end,
                transaction.source_file,
                transaction.source_page,
                transaction.extraction_method,
                transaction.source_row,
            ]
        )

    for row in range(2, sheet.max_row + 1):
        for column in (2, 4, 10, 11, 13, 14):
            sheet.cell(row, column).number_format = "@"
        sheet.cell(row, 3).number_format = "yyyy-mm-dd hh:mm:ss"
        for column in (7, 8, 9):
            sheet.cell(row, column).number_format = "#,##0.00"
        for column in (15, 16):
            sheet.cell(row, column).number_format = "yyyy-mm-dd"
        for column in range(1, len(DETAIL_HEADERS) + 1):
            sheet.cell(row, column).alignment = Alignment(vertical="top", wrap_text=True)

    _apply_header(sheet, len(DETAIL_HEADERS))
    _fit_columns(sheet)


def _write_summary(sheet, batch: BatchResult) -> None:
    sheet.append(SUMMARY_HEADERS)
    for statement in batch.statements:
        extraction = statement.extraction
        validation = statement.validation
        balance_status = (
            "通过"
            if not validation.balance_break_rows
            else "失败（第" + "、".join(str(row) for row in validation.balance_break_rows) + "笔）"
        )
        sheet.append(
            [
                statement.source_path.name,
                validation.status,
                extraction.bank_name if extraction else "",
                extraction.account_number if extraction else "",
                extraction.period_start if extraction else None,
                extraction.period_end if extraction else None,
                extraction.page_count if extraction else None,
                extraction.declared_count if extraction else None,
                validation.actual_count,
                extraction.declared_debit if extraction else None,
                validation.actual_debit,
                extraction.declared_credit if extraction else None,
                validation.actual_credit,
                balance_status,
                extraction.extraction_method if extraction else "",
                "；".join(validation.errors),
                extraction.currency if extraction else "",
                extraction.opening_balance if extraction else None,
                extraction.opening_page if extraction else None,
                extraction.closing_balance if extraction else None,
                extraction.closing_page if extraction else None,
                extraction.statement_start if extraction else None,
                extraction.statement_end if extraction else None,
                *[value for name in ("期初期末", "首笔余额", "末笔余额")
                  for value in (validation.control_checks.get(name, {}).get("status", "未检查") +
                                ("，" + validation.control_checks[name]["reason"] if validation.control_checks.get(name, {}).get("reason") else ""),
                                validation.control_checks.get(name, {}).get("difference"))],
                str(statement.source_path),
                validation.validation_basis,
                "；".join(validation.warnings),
            ]
        )
        row = sheet.max_row
        if validation.status == "核对失败":
            for cell in sheet[row]:
                cell.fill = FAILED_FILL
                cell.font = FAILED_FONT
        elif validation.status == "零交易通过":
            for cell in sheet[row]:
                cell.fill = ZERO_FILL
                cell.font = ZERO_FONT

    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, 4).number_format = "@"
        for column in (5, 6):
            sheet.cell(row, column).number_format = "yyyy-mm-dd"
        for column in (18, 20, 25, 27, 29):
            sheet.cell(row, column).number_format = "#,##0.00"
        for column in (22, 23):
            sheet.cell(row, column).number_format = "yyyy-mm-dd"
        for column in (10, 11, 12, 13):
            sheet.cell(row, column).number_format = "#,##0.00"
        for column in range(1, len(SUMMARY_HEADERS) + 1):
            sheet.cell(row, column).alignment = Alignment(vertical="top", wrap_text=True)

    _apply_header(sheet, len(SUMMARY_HEADERS))
    _fit_columns(sheet, maximum=55)


def write_batch_workbook(
    batch: BatchResult,
    source_root: Path,
    now: datetime | None = None,
) -> Path:
    output_path = _unique_output_path(source_root, now or datetime.now())
    workbook = Workbook()
    detail_sheet = workbook.active
    detail_sheet.title = "流水明细"
    summary_sheet = workbook.create_sheet("提取汇总")
    _write_details(detail_sheet, batch)
    _write_summary(summary_sheet, batch)
    links = workbook.create_sheet("期间衔接检查")
    links.append(LINK_HEADERS)
    for record in statement_links(batch):
        links.append([record[key] for key in LINK_HEADERS])
    for row in links.iter_rows(min_row=2):
        row[1].number_format = "@"
        for cell in row[5:9]:
            cell.number_format = "yyyy-mm-dd"
        for cell in row[9:12]:
            cell.number_format = "#,##0.00"
    _apply_header(links, len(LINK_HEADERS))
    _fit_columns(links, maximum=55)
    pending = [tx for s in batch.statements if s.extraction is not None
               and s.validation.status == "核对失败"
               for tx in s.extraction.transactions]
    if pending:
        _write_details(workbook.create_sheet("待核查流水"), batch, pending)
    workbook.save(output_path)

    checked = load_workbook(output_path, read_only=True, data_only=False)
    try:
        if checked.sheetnames[:3] != ["流水明细", "提取汇总", "期间衔接检查"]:
            raise ValueError("生成的工作簿工作表不完整")
    finally:
        checked.close()
    return output_path
